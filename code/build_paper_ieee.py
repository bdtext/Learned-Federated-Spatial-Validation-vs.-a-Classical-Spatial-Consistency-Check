"""Build the IEEE Access-formatted .docx from paper/paper.md and paper/authors.json.
Layout follows the IEEE Access article template: US Letter, 0.625-in side margins, single-column front matter (publication
lines, title, author line with superscripts, affiliations, corresponding author, funding, ABSTRACT, INDEX TERMS), two-column
Times New Roman 10 pt body with sans-serif blue section headings ("I. INTRODUCTION") and bold italic uppercase subsections
("A. ..."), table captions above tables ("TABLE 1."), figure captions below figures ("FIGURE 1."), wide tables and figures
spanning both columns, running headers and VOLUME footers, 8 pt references, author biographies with photographs below the
references, and supplementary tables at the end.
Usage: python build_paper_ieee.py ../paper/paper.md ../paper/paper_ieee_access.docx
"""
import os, re, sys, json
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.join(os.path.dirname(__file__), "..")
BLUE = RGBColor(0x1C, 0x5B, 0xA9)
SANS, SERIF = "Arial", "Times New Roman"
PAGE_W, MARGIN_LR = 8.5, 0.625
TEXT_W = PAGE_W - 2 * MARGIN_LR          # 7.25 in
COL_GAP = 0.25
COL_W = (TEXT_W - COL_GAP) / 2           # 3.5 in
JUST, LEFT, CENTER, RIGHT = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT

# ---------------------------------------------------------------- low-level helpers
def set_cols(section, n, gap_in=COL_GAP):
    sectPr = section._sectPr
    for c in sectPr.findall(qn("w:cols")):
        sectPr.remove(c)
    cols = OxmlElement("w:cols"); cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(int(gap_in * 1440)))
    sectPr.append(cols)

def new_section(doc, n_cols):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(sec, n_cols)
    return sec

def style_run(run, font=SERIF, size=10, bold=False, italic=False, color=None, superscript=False, small_caps=False):
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    rpr = run._element.get_or_add_rPr(); rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), font)
    if color is not None: run.font.color.rgb = color
    if superscript: run.font.superscript = True
    if small_caps: run.font.small_caps = True
    return run

def typeset(t):
    """Journal typography for the plain-text source: +/- -> ±, degC -> °C, sigma_a -> σ with subscript, 10-3 -> 10 with superscript."""
    t = t.replace("+/-", "±").replace("degC", "°C")
    t = re.sub(r"\bsigma_a\b", "σ_{a}", t); t = re.sub(r"\bsigma_hf\b", "σ_{hf}", t)
    t = re.sub(r"(\d)x10-(\d)\b", r"\1×10^{-\2}", t); t = re.sub(r"\b10-(\d)\b", r"10^{-\1}", t)
    t = t.replace("m s-1", "m s^{-1}").replace(" x ", " × ")
    return t

def add_runs(par, text, font=SERIF, size=10, bold=False, italic=False, color=None):
    """Inline markup: **bold**, *italic*, `code`, ^{superscript}, _{subscript}."""
    text = typeset(text); pos = 0
    for m in re.finditer(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^\{[^}]+\}|_\{[^}]+\})", text):
        if m.start() > pos: style_run(par.add_run(text[pos:m.start()]), font, size, bold, italic, color)
        t = m.group(0)
        if t.startswith("**"): style_run(par.add_run(t[2:-2]), font, size, True, italic, color)
        elif t.startswith("`"): style_run(par.add_run(t[1:-1]), "Consolas", size - 1, bold, italic, color)
        elif t.startswith("^{"): style_run(par.add_run(t[2:-1]), font, size, bold, italic, color, superscript=True)
        elif t.startswith("_{"): r = style_run(par.add_run(t[2:-1]), font, size, bold, italic, color); r.font.subscript = True
        else: style_run(par.add_run(t[1:-1]), font, size, bold, True, color)
        pos = m.end()
    if pos < len(text): style_run(par.add_run(text[pos:]), font, size, bold, italic, color)

def para(doc, text="", font=SERIF, size=10, bold=False, italic=False, color=None, align=JUST, before=0, after=0, first_indent=None,
         left_indent=None, hanging=None, keep_next=False, line=1.0):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if first_indent is not None: pf.first_line_indent = Inches(first_indent)
    if left_indent is not None: pf.left_indent = Inches(left_indent)
    if hanging is not None: pf.left_indent = Inches(hanging); pf.first_line_indent = Inches(-hanging)
    if keep_next: pf.keep_with_next = True
    if text: add_runs(p, text, font, size, bold, italic, color)
    return p

def add_field(par, instr, font=SANS, size=7.5):
    r = par.add_run(); style_run(r, font, size)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._element.append(f1); r._element.append(it); r._element.append(f2)

def wordmark(par, size=16):
    style_run(par.add_run("IEEE"), SANS, size, bold=True, color=BLUE); style_run(par.add_run("Access"), SANS, size, italic=True, color=BLUE)

def set_cell_borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)

def fmt(x):
    """Numbers from numeric CSV columns: at most three decimals, trailing zeros removed; strings are kept as written."""
    if isinstance(x, float):
        if x != x: return "n/a"
        s = f"{x:.3f}".rstrip("0").rstrip("."); return s if s not in ("", "-0") else "0"
    s = str(x)
    return "n/a" if s == "" else s

# ---------------------------------------------------------------- content builders
def caption_split(cap):
    m = re.match(r"^(Table|Fig\.|Figure)\s*(S?\d+[a-z]?)\.\s*(.*)$", cap.strip())
    if m:
        kind = "TABLE" if m.group(1) == "Table" else "FIGURE"
        return kind, m.group(2), m.group(3)
    return None, None, cap

def add_table(doc, rows, caption, wide):
    kind, num, text = caption_split(caption)
    p = para(doc, align=LEFT, before=6, after=2, keep_next=True)
    style_run(p.add_run(f"{kind or 'TABLE'} {num or ''}. ".replace(" .", ".")), SANS, 8, bold=True)
    add_runs(p, text, SANS, 8, bold=True)
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.CENTER; set_cell_borders(t); t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j); cell.text = ""
            cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.space_before = Pt(0)
            add_runs(cp, str(row[j]) if j < len(row) else "", SERIF, 7.5 if wide else 7, bold=(i == 0))
    para(doc, after=4)

def add_figure(doc, path, caption, width_in):
    kind, num, text = caption_split(caption)
    p = doc.add_paragraph(); p.paragraph_format.alignment = CENTER; p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_in))
    c = para(doc, align=LEFT, before=2, after=8)
    style_run(c.add_run(f"{kind or 'FIGURE'} {num or ''}. ".replace(" .", ".")), SANS, 8, bold=True)
    add_runs(c, text, SANS, 8)

def heading(doc, text, level):
    text = re.sub(r"\s+", " ", text.strip())
    if level == 2:
        p = para(doc, align=LEFT, before=10, after=3, keep_next=True); style_run(p.add_run(text.upper()), SANS, 10, bold=True, color=BLUE)
    elif level == 3:
        p = para(doc, align=LEFT, before=8, after=2, keep_next=True); style_run(p.add_run(text.upper()), SANS, 9, bold=True, italic=True)
    else:
        p = para(doc, align=LEFT, before=6, after=2, keep_next=True); style_run(p.add_run(text), SANS, 9, italic=True)

# ---------------------------------------------------------------- document assembly
def main(src, dst):
    A = json.load(open(os.path.join(ROOT, "paper", "authors.json"), encoding="utf-8"))
    lines = open(src, encoding="utf-8").read().splitlines()
    # ---- parse the markdown into blocks
    title = next(l[2:].strip() for l in lines if l.startswith("# "))
    def section(name):
        i = next(k for k, l in enumerate(lines) if l.strip() == f"## {name}")
        j = next((k for k in range(i + 1, len(lines)) if lines[k].startswith("## ")), len(lines))
        return lines[i + 1:j]
    abstract = " ".join(l.strip() for l in section("Abstract") if l.strip() and not l.startswith("**Index terms"))
    index_terms = next(l for l in lines if l.startswith("**Index terms:**")).split("**", 2)[2].strip()
    body_start = next(k for k, l in enumerate(lines) if l.startswith("## I. "))
    refs_start = next(k for k, l in enumerate(lines) if l.strip() == "## References")
    supp_start = next((k for k, l in enumerate(lines) if l.strip().startswith("## Supplementary")), len(lines))
    body = lines[body_start:refs_start]; refs = [l for l in lines[refs_start + 1:supp_start] if l.strip().startswith("[")]
    supp = lines[supp_start + 1:] if supp_start < len(lines) else []
    # tables destined for the supplement (captions "Table S...") are pulled out of the body
    supp_blocks, body_kept = [], []
    for l in body:
        m = re.match(r"\{\{csv:([^|]+)\|([^|]*)\|([^}]*)\}\}", l.strip()); mf = re.match(r"\{\{fig:([^|]+)\|([^|]*)\|?([^}]*)\}\}", l.strip())
        if (m and m.group(3).strip().startswith("Table S")) or (mf and mf.group(2).strip().startswith("Fig. S")):
            supp_blocks.append(l)
        else:
            body_kept.append(l)
    supp_blocks += [l for l in supp if l.strip().startswith("{{")]
    index_terms = index_terms.replace("; ", ", ")

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = SERIF; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    sec = doc.sections[0]
    sec.page_width = Inches(PAGE_W); sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(MARGIN_LR); sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.9)
    sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.4)
    sec.different_first_page_header_footer = True
    # headers and footers
    def two_cell(container, left_fn, right_fn):
        """Borderless one-row table: left cell content, right cell right-aligned content (robust header/footer alignment)."""
        t = container.add_table(rows=1, cols=2, width=Inches(TEXT_W)); t.autofit = False
        widths = (1.6, TEXT_W - 1.6)
        t.columns[0].width = Inches(widths[0]); t.columns[1].width = Inches(widths[1])
        for c, w in zip(t.rows[0].cells, widths):
            c.width = Inches(w); c.text = ""
        pl, pr = t.cell(0, 0).paragraphs[0], t.cell(0, 1).paragraphs[0]
        pl.alignment = LEFT; pr.alignment = RIGHT; left_fn(pl); right_fn(pr)
        return t
    h = sec.first_page_header.paragraphs[0]; h.alignment = RIGHT; wordmark(h, 16)
    sec.header.paragraphs[0].text = ""
    two_cell(sec.header, lambda p: wordmark(p, 12), lambda p: style_run(p.add_run(A["running_head"]), SANS, 7.5))
    for footer in (sec.first_page_footer, sec.footer):
        footer.paragraphs[0].text = ""
        two_cell(footer, lambda p: style_run(p.add_run(A["volume_line"]), SANS, 7.5), lambda p: add_field(p, "PAGE"))
    # ---- front matter (single column)
    para(doc, "Date of publication xxxx 00, 0000, date of current version xxxx 00, 0000.", SERIF, 8, align=LEFT, after=2)
    para(doc, A["doi_line"], SERIF, 7.5, italic=True, align=LEFT, after=14)
    p = para(doc, align=LEFT, after=10, line=1.0); style_run(p.add_run(title), SANS, 22, bold=True, color=BLUE)
    p = para(doc, align=LEFT, after=2)
    auth = A["authors"]
    for i, a in enumerate(auth):
        sep = "" if i == 0 else (", AND " if i == len(auth) - 1 else ", ")
        if sep: style_run(p.add_run(sep), SANS, 10, bold=True)
        style_run(p.add_run(a["name"].upper()), SANS, 10, bold=True)
        style_run(p.add_run(",".join(str(x) for x in a["affil"])), SANS, 10, bold=True, superscript=True)
    for k, v in A["affiliations"].items():
        p = para(doc, align=LEFT, after=0); style_run(p.add_run(k), SERIF, 7.5, superscript=True); style_run(p.add_run(v), SERIF, 7.5)
    para(doc, A["corresponding"], SERIF, 8, align=LEFT, before=6, after=6)
    para(doc, A["funding"], SERIF, 8, align=LEFT, after=14)
    p = para(doc, align=JUST, after=10, left_indent=0.15); style_run(p.add_run("ABSTRACT "), SANS, 10, bold=True, color=BLUE); add_runs(p, abstract, SERIF, 10)
    p = para(doc, align=JUST, after=14, left_indent=0.15); style_run(p.add_run("INDEX TERMS "), SANS, 10, bold=True, color=BLUE); add_runs(p, index_terms, SERIF, 10)
    # ---- body (two columns)
    new_section(doc, 2)
    i = 0; after_heading = True
    while i < len(body_kept):
        line = body_kept[i]
        m = re.match(r"\{\{csv:([^|]+)\|([^|]*)\|([^}]*)\}\}", line.strip())
        if m:
            path, cols, cap = m.group(1), m.group(2), m.group(3)
            df = pd.read_csv(os.path.join(ROOT, path), keep_default_na=False, dtype=str)  # keep the numbers exactly as the tables wrote them
            if cols.strip(): df = df[[c.strip() for c in cols.split(",")]]
            rows = [list(df.columns)] + [[fmt(v) for v in r] for r in df.itertuples(index=False)]
            wide = len(df.columns) >= 6 or max(len(str(v)) for r in rows for v in r) > 28
            if wide: new_section(doc, 1)
            add_table(doc, rows, cap, wide)
            if wide: new_section(doc, 2)
            i += 1; after_heading = False; continue
        m = re.match(r"\{\{fig:([^|]+)\|([^|]*)\|?([^}]*)\}\}", line.strip())
        if m:
            path, cap, w = m.group(1), m.group(2), m.group(3)
            p = os.path.join(ROOT, path)
            if os.path.exists(p):
                w = float(w) if w else 3.4
                wide = w >= 6.0
                if wide: new_section(doc, 1)
                add_figure(doc, p, cap, min(w, TEXT_W) if wide else min(w, COL_W))
                if wide: new_section(doc, 2)
            i += 1; after_heading = False; continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#")); heading(doc, line.lstrip("#").strip(), level); after_heading = True
        elif line.startswith("- "):
            p = para(doc, align=JUST, hanging=0.15, after=2); add_runs(p, "• " + line[2:], SERIF, 10); after_heading = False
        elif re.match(r"^\d+\. ", line):
            para(doc, line, SERIF, 10, hanging=0.2, after=2); after_heading = False
        elif line.strip() == "":
            pass
        else:
            para(doc, line, SERIF, 10, first_indent=None if after_heading else 0.15); after_heading = False
        i += 1
    # ---- acknowledgment, references, biographies (two columns)
    heading(doc, "Acknowledgment", 2); para(doc, A["acknowledgment"], SERIF, 9)
    heading(doc, "References", 2)
    for r in refs:
        m = re.match(r"^(\[\d+\])\s*(.*)$", r.strip())
        p = para(doc, align=JUST, after=1, hanging=0.3); style_run(p.add_run(m.group(1) + " "), SERIF, 8); add_runs(p, m.group(2), SERIF, 8)
    para(doc, after=6)
    for a in auth:
        bio = A["biographies"].get(a["name"], "[Biography to be supplied.]")
        photo = os.path.join(ROOT, "paper", a.get("photo", "")) if a.get("photo") else ""
        t = doc.add_table(rows=1, cols=2); t.autofit = False
        t.columns[0].width = Inches(1.05); t.columns[1].width = Inches(COL_W - 1.1)
        c0, c1 = t.cell(0, 0), t.cell(0, 1); c0.width = Inches(1.05); c1.width = Inches(COL_W - 1.1)
        c0.text = ""; c1.text = ""
        if photo and os.path.exists(photo):
            c0.paragraphs[0].add_run().add_picture(photo, width=Inches(0.95))
        cp = c1.paragraphs[0]; cp.paragraph_format.alignment = JUST
        style_run(cp.add_run(a["name"].upper() + " "), SANS, 8, bold=True); add_runs(cp, bio, SERIF, 8)
        para(doc, after=8)
    # ---- supplementary tables (full width)
    if supp_blocks:
        sec_s = doc.add_section(WD_SECTION.NEW_PAGE); set_cols(sec_s, 1)
        heading(doc, "Supplementary Material", 2)
        para(doc, "The supplementary tables and figures referenced from the main text: labelled fractions, injection audit, trained validators, block-bootstrap intervals, per-variable, per-family and event-level slices of the detection benchmark, the per-family slice of the subtle regime, realized fault magnitudes, the complete set of paired contrasts, per-seed differences and learning curves.", SERIF, 9, after=6)
        for l in supp_blocks:
            m = re.match(r"\{\{csv:([^|]+)\|([^|]*)\|([^}]*)\}\}", l.strip())
            if m:
                path, cols, cap = m.group(1), m.group(2), m.group(3)
                if not os.path.exists(os.path.join(ROOT, path)): continue
                df = pd.read_csv(os.path.join(ROOT, path), keep_default_na=False, dtype=str)
                if cols.strip(): df = df[[c.strip() for c in cols.split(",")]]
                rows = [list(df.columns)] + [[fmt(v) for v in r] for r in df.itertuples(index=False)]
                add_table(doc, rows, cap, True); continue
            mf = re.match(r"\{\{fig:([^|]+)\|([^|]*)\|?([^}]*)\}\}", l.strip())
            if mf and os.path.exists(os.path.join(ROOT, mf.group(1))):
                add_figure(doc, os.path.join(ROOT, mf.group(1)), mf.group(2), min(float(mf.group(3)) if mf.group(3) else 4.5, 5.0))
    doc.save(dst); print("wrote", dst)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
